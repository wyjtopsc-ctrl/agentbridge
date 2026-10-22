#!/usr/bin/env python3
"""Render a markdown strategy/plan document into a styled DOCX.

  python3 tools/make_memo.py --source docs/x.md --title "..." --subtitle "..." \
      --version 1.0 --output deliverables/x.docx
"""
import argparse
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import docx_kit  # noqa: E402
from docx import Document  # noqa: E402


def parse(markdown: str):
    blocks, lines, i = [], markdown.splitlines(), 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("### "):
            blocks.append(("h2", line[4:]))
        elif line.startswith("## "):
            blocks.append(("h1", line[3:]))
        elif line.startswith("# "):
            pass
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip().strip("`") for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            if rows:
                blocks.append(("table", rows[0], rows[1:]))
            continue
        elif line.startswith("> "):
            blocks.append(("quote", line[2:]))
        elif line.startswith("- "):
            blocks.append(("bullet", line[2:]))
        elif re.match(r"^\d+\.\s", line):
            blocks.append(("num", re.sub(r"^\d+\.\s", "", line)))
        else:
            blocks.append(("p", line.replace("`", "")))
        i += 1
    return blocks


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--source", required=True)
    ap.add_argument("--title", required=True)
    ap.add_argument("--subtitle", default="")
    ap.add_argument("--version", default="1.0")
    ap.add_argument("--author", default="agentbridge")
    ap.add_argument("--output", required=True)
    args = ap.parse_args()

    blocks = parse(open(args.source, encoding="utf-8").read())
    doc = docx_kit.new_document(args.title, args.subtitle, args.version, author=args.author)
    docx_kit.render(doc, blocks)
    os.makedirs(os.path.dirname(os.path.abspath(args.output)), exist_ok=True)
    docx_kit.save(doc, args.output)

    reopened = Document(args.output)
    print(f"{args.output}\n  blocks={len(blocks)} paragraphs={len(reopened.paragraphs)} "
          f"tables={len(reopened.tables)}")
    assert len(reopened.paragraphs) > 10, "document looks empty"


if __name__ == "__main__":
    main()
