#!/usr/bin/env python3
"""Build the Word deliverables from the content modules."""
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

from docx import Document  # noqa: E402

import docx_kit  # noqa: E402
import manual_content  # noqa: E402
import whitepaper_content  # noqa: E402

OUT = os.path.join(os.path.dirname(HERE), "deliverables")


def product_version() -> str:
    import re
    text = open(os.path.join(os.path.dirname(HERE), "src", "agentbridge", "__init__.py"),
                encoding="utf-8").read()
    return re.search(r'__version__\s*=\s*"([^"]+)"', text).group(1)


VERSION = product_version()


def build(content, filename: str) -> str:
    content.VERSION = VERSION
    doc = docx_kit.new_document(content.TITLE, content.SUBTITLE, content.VERSION)
    docx_kit.render(doc, content.BLOCKS)
    path = os.path.join(OUT, filename)
    docx_kit.save(doc, path)

    reopened = Document(path)
    words = sum(len(p.text.split()) for p in reopened.paragraphs)
    print(f"{path}\n  paragraphs={len(reopened.paragraphs)} tables={len(reopened.tables)} "
          f"chars={sum(len(p.text) for p in reopened.paragraphs)}")
    assert len(reopened.paragraphs) > 40, "document looks empty"
    assert len(reopened.tables) >= 5, "tables missing"
    return path


if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    build(whitepaper_content, f"agentbridge-whitepaper-{VERSION}.docx")
    build(manual_content, f"agentbridge-manual-{VERSION}.docx")
